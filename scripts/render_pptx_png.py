"""Approximate PPTX -> PNG renderer for visual QA.

LibreOffice headless conversion is unavailable in some sandboxes, so this draws
each slide directly from python-pptx: solid-filled rectangles plus textboxes
(runs with size/bold/color/align, vertical anchor, line spacing, char-level word
wrap). It is not pixel-perfect — fonts differ and mono fonts here lack Japanese
glyphs (real PowerPoint/LibreOffice falls back to a JP font) — but it is faithful
enough to judge layout density, whitespace and alignment while iterating on the
deck design.

Usage:
    python render_pptx_png.py deck.pptx outdir [slide_number ...]
"""
import sys
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

EMU = 914400
DPI = 120
JP     = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
MONO   = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
MONO_B = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf"

_fonts = {}


def font(path, px):
    key = (path, px)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(path, px)
    return _fonts[key]


def px(emu):
    return int(round(emu / EMU * DPI))


def pt_px(pt):
    return int(round(pt / 72.0 * DPI))


def is_mono(name):
    return bool(name) and ("Consolas" in name or "mono" in name.lower())


def width(draw, text, fnt):
    b = draw.textbbox((0, 0), text, font=fnt)
    return b[2] - b[0]


def wrap(draw, runs, maxw):
    """runs: [(text, fnt, color)] -> list of lines, each [(text, fnt, color)]."""
    lines = [[]]
    cur = 0
    for text, fnt, col in runs:
        tokens, buf = [], ""
        for ch in text:
            if ch == "\n":
                if buf:
                    tokens.append(buf); buf = ""
                tokens.append("\n")
            elif ord(ch) < 128 and ch != " ":   # keep ascii words intact
                buf += ch
            else:
                if buf:
                    tokens.append(buf); buf = ""
                tokens.append(ch)
        if buf:
            tokens.append(buf)
        for tok in tokens:
            if tok == "\n":
                lines.append([]); cur = 0; continue
            w = width(draw, tok, fnt)
            if cur + w > maxw and cur > 0:
                lines.append([]); cur = 0
            lines[-1].append((tok, fnt, col)); cur += w
    return lines


def run_font(r):
    sz = r.font.size.pt if r.font.size else 12
    name = r.font.name or "Noto Sans JP"
    if is_mono(name):
        return font(MONO_B if r.font.bold else MONO, pt_px(sz)), sz
    return font(JP, pt_px(sz)), sz


def run_color(r):
    try:
        if r.font.color and r.font.color.type is not None:
            c = r.font.color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return (20, 20, 20)


def render_slide(prs, slide, out):
    W, H = px(prs.slide_width), px(prs.slide_height)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    for shp in slide.shapes:
        try:
            L, T, Wd, Ht = px(shp.left), px(shp.top), px(shp.width), px(shp.height)
        except Exception:
            continue
        try:
            if shp.fill.type == 1:  # solid
                c = shp.fill.fore_color.rgb
                d.rectangle([L, T, L + Wd, T + Ht], fill=(c[0], c[1], c[2]))
        except Exception:
            pass
        if not shp.has_text_frame:
            continue
        tf = shp.text_frame
        blocks = []          # (wrapped_lines, line_h, space_after_px, align)
        total_h = 0
        for para in tf.paragraphs:
            runs, maxsz = [], 12
            for r in para.runs:
                fnt, sz = run_font(r)
                maxsz = max(maxsz, sz)
                runs.append((r.text, fnt, run_color(r)))
            ls = para.line_spacing if para.line_spacing and para.line_spacing > 0 else 1.1
            lh = pt_px(maxsz) * ls
            sa = pt_px(para.space_after.pt) if para.space_after else 0
            lines = wrap(d, runs, Wd) if runs else [[]]
            blocks.append((lines, lh, sa, para.alignment))
            total_h += lh * len(lines) + sa
        if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
            y = T + max(0, (Ht - total_h) / 2)
        elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
            y = T + max(0, Ht - total_h)
        else:
            y = T
        for lines, lh, sa, align in blocks:
            for line in lines:
                lw = sum(width(d, t, f) for (t, f, _) in line)
                if align == PP_ALIGN.CENTER:
                    x = L + (Wd - lw) / 2
                elif align == PP_ALIGN.RIGHT:
                    x = L + (Wd - lw)
                else:
                    x = L
                for (t, f, c) in line:
                    d.text((x, y), t, font=f, fill=c)
                    x += width(d, t, f)
                y += lh
            y += sa
    img.save(out)


def main():
    src, outdir = sys.argv[1], sys.argv[2]
    Path(outdir).mkdir(parents=True, exist_ok=True)
    prs = Presentation(src)
    want = {int(x) - 1 for x in sys.argv[3:]}
    for i, slide in enumerate(prs.slides):
        if want and i not in want:
            continue
        out = Path(outdir) / f"slide{i + 1}.png"
        render_slide(prs, slide, out)
        print("rendered", out)


if __name__ == "__main__":
    main()
