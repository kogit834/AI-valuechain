"""ANDPAD 本部総会 向けプレゼン資料（本編5分・7枚＋付録）を生成する。
デザイン方針: 箱の多用を避けつつ、大きめの文字・詰めた余白・細い罫線・赤の縦
アクセントで“メリハリのあるエディトリアル調”。赤はアクセントとデータ強調に限定。
付録: data/companies.csv・segments.csv からバリューチェーン別キー企業一覧（全社）を
自動生成し、本編の後ろに付ける。
出力: output/andpad_soukai_ai_deck.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pathlib import Path
import csv

# ---- palette ----
RED   = RGBColor(0xE6,0x00,0x12)
RED2  = RGBColor(0xF0,0x66,0x70)   # 淡い赤（同系2トーン）
PINK  = RGBColor(0xFF,0xCF,0xD3)   # 赤地の上の小ラベル用
INK   = RGBColor(0x15,0x18,0x1C)
SUB   = RGBColor(0x37,0x3C,0x43)
MUTE  = RGBColor(0x6E,0x76,0x7F)
FAINT = RGBColor(0x9A,0xA1,0xA9)
HAIR  = RGBColor(0xD6,0xDA,0xDF)
G1    = RGBColor(0xD2,0xD6,0xDB)   # バー淡グレー
G_DK  = RGBColor(0x53,0x59,0x61)   # バー濃グレー
WHITE = RGBColor(0xFF,0xFF,0xFF)

FONT   = "Noto Sans JP"
FONT_M = "Consolas"

LM, RM = 0.9, 12.43
CW = RM - LM   # 11.53

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = WHITE; r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s


def box(s, x, y, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    return sh


def hrule(s, x, y, w, color=HAIR, wt=1.0):
    box(s, x, y, w, wt/72.0, color)


def vrule(s, x, y, h, color=HAIR, wt=1.0):
    box(s, x, y, wt/72.0, h, color)


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line_spacing=1.1):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_before = Pt(0); p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, font, tracking) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = font
            rPr = r._r.get_or_add_rPr()
            ea = rPr.find(qn('a:ea'))
            if ea is None:
                ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
            ea.set('typeface', font)
            if tracking: rPr.set('spc', str(int(tracking * 100)))
    return tb


def R(txt, size, color=INK, bold=False, font=FONT, tracking=0):
    return (txt, size, color, bold, font, tracking)


def takeaway(s, y, runs, h=0.6):
    """赤い縦アクセントバー＋大きめテキスト（塗りなしでメリハリ）。"""
    box(s, LM, y+0.05, 0.09, h-0.1, RED)
    text(s, LM+0.28, y, CW-0.3, h, runs, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.16)


def header(s, page, eyebrow, title):
    box(s, LM, 0.62, 0.14, 0.14, RED)
    text(s, 1.14, 0.55, 9.5, 0.28, [[R(eyebrow, 11.5, RED, True, FONT_M, 1.5)]])
    text(s, RM-0.7, 0.55, 0.7, 0.28, [[R(page, 11.5, FAINT, True, FONT_M)]], align=PP_ALIGN.RIGHT)
    text(s, LM-0.02, 0.92, CW, 0.55, [[R(title, 24, INK, True)]])
    hrule(s, LM, 1.6, CW, wt=1.4)


# =========================================================
# SLIDE 1 — 表紙
# =========================================================
s = slide()
box(s, LM, 1.6, 0.15, 0.15, RED)
text(s, 1.16, 1.52, 10.5, 0.3,
     [[R("ANDPAD 本部総会   ·   WHY MANUFACTURING NOW", 13, RED, True, FONT_M, 2)]])
text(s, LM-0.04, 2.45, 11.8, 2.2,
     [[R("建設SaaSの我々が、", 46, INK, True)],
      [R("なぜ、いま", 46, INK, True), R("製造業", 46, RED, True), R("なのか。", 46, INK, True)]],
     line_spacing=1.12)
box(s, LM, 4.85, 3.2, 0.05, RED)
text(s, LM-0.02, 5.15, 11.7, 1.1,
     [[R("AIをはじめ世界のメガトレンドは、膨大な“設備”を必要とする。", 17, SUB)],
      [R("それを作り・建てるのは、建設業だけでなく製造業がコア——その市場が、いま圧倒的に広がっている。", 17, SUB)]],
     line_spacing=1.5)
hrule(s, LM, 6.95, CW)
text(s, LM-0.02, 7.06, 11.7, 0.3,
     [[R("2026-07   /   AIインフラ・バリューチェーン調査（国内外154社・決算/IR一次情報）＋公開データ", 10.5, FAINT, False, FONT_M)]])

# =========================================================
# SLIDE 2 — メガトレンド → 設備 → 製造業がコア → AIにフォーカス
# =========================================================
s = slide()
header(s, "01", "世界のメガトレンドと、その“共通の土台”",
       "世界が動く先には、必ず「設備」がいる。その主役は、製造業だ。")

text(s, LM-0.02, 1.72, CW, 0.3,
     [[R("いま世界を動かす3つのメガトレンド。どれも、実現するには物理的な“設備”が要る。", 13.5, MUTE, True)]])

cols = [
    ("01", "AI", True,  "生成AI・データセンター・半導体", "電源・冷却・建屋・チップ工場"),
    ("02", "脱炭素（GX）", False, "再エネ・送電網・電化・蓄電池", "発電・変電・ケーブル・電池／EV工場"),
    ("03", "経済安保・国内回帰", False, "半導体・重要物資の国産化", "国内に工場を新設・供給網を再構築"),
]
pitch = CW / 3
ctop = 2.24
for i, (no, nm, hot, what, need) in enumerate(cols):
    x = LM + i * pitch
    w = pitch - 0.44
    if i > 0:
        vrule(s, x-0.24, ctop+0.02, 1.98, wt=1.0)
    box(s, x, ctop, w, 0.07, RED if hot else G_DK)
    text(s, x, ctop+0.22, w, 0.24, [[R(no, 11, FAINT, True, FONT_M, 1)]])
    if hot:
        text(s, x, ctop+0.2, w, 0.24, [[R("● 本日フォーカス", 10, RED, True, FONT_M)]],
             align=PP_ALIGN.RIGHT)
    text(s, x, ctop+0.5, w, 0.5, [[R(nm, 22, (RED if hot else INK), True)]])
    text(s, x, ctop+1.06, w, 0.3, [[R(what, 13, MUTE, True)]], line_spacing=1.16)
    text(s, x, ctop+1.44, w, 0.6,
         [[R("要る設備", 9.5, FAINT, True, FONT_M, 1)],
          [R("→ ", 13.5, RED, True), R(need, 13.5, SUB, True)]],
         line_spacing=1.2, space_after=2)

# 3列 → ひとつの土台へ収束（▼）
for i in range(3):
    x = LM + i * pitch
    text(s, x, 4.34, pitch-0.44, 0.3, [[R("▼", 13, RED, True)]], align=PP_ALIGN.CENTER)

# 収束先：赤の全幅バンド（P5と同系の“塗り”でメッセージを運ぶ）
by, bh = 4.66, 1.12
box(s, LM, by, CW, bh, RED)
split = CW * 0.46
box(s, LM+split-0.01, by+0.18, 0.02, bh-0.36, WHITE)
text(s, LM+0.42, by, split-0.6, bh,
     [[R("3つに共通する“土台”", 10.5, PINK, True, FONT_M, 1)],
      [R("膨大な“設備”がいる。", 23, WHITE, True)]],
     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1, space_after=3)
text(s, LM+split+0.5, by, CW-split-0.9, bh,
     [[R("作る・建てる主役は", 10.5, PINK, True, FONT_M, 1)],
      [R("重電・機械・電機など、“製造業”がコア。", 18, WHITE, True)]],
     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.14, space_after=3)

takeaway(s, 6.18,
         [[R("→ 今日はその中で、最も勢いのある「", 19, INK, True),
           R("AI", 19, RED, True), R("」に絞って話す。", 19, INK, True)]], h=0.62)

# =========================================================
# SLIDE 3 — トレンド：これまでの何倍か
# =========================================================
s = slide()
header(s, "02", "規模より“勢い”——トレンドとして、どれだけ急か",
       "AI投資は“大きい”だけじゃない。数年で“何倍”に跳ねている。")


def vbars(s, x0, base_y, area_w, max_h, values, years, val_label):
    n = len(values); mx = max(values)
    slot = area_w / n; bw = min(0.5, slot * 0.56)
    hrule(s, x0, base_y, area_w, color=HAIR, wt=1.4)
    for i, v in enumerate(values):
        h = max_h * v / mx
        bxx = x0 + i * slot + (slot - bw) / 2
        box(s, bxx, base_y - h, bw, h, RED if i == n-1 else G1)
        text(s, x0 + i*slot, base_y + 0.07, slot, 0.24,
             [[R(years[i], 9.5, MUTE, True, FONT_M)]], align=PP_ALIGN.CENTER)
        if i == n-1:
            text(s, bxx-0.6, base_y-h-0.26, bw+1.2, 0.24,
                 [[R(val_label, 10, MUTE, True, FONT_M)]], align=PP_ALIGN.CENTER)


rows = [
    ("AIチップの需要", "NVIDIA データセンター売上", [15, 48, 115, 196],
     ["FY23", "FY24", "FY25", "FY26"], "約2,000億ドル", "約13倍"),
    ("AIインフラ投資", "Big Tech 4社の設備投資（年間）", [180, 230, 410, 725],
     ["’23", "’24", "’25", "’26"], "7,250億ドル", "約4倍"),
    ("国内DC建設投資", "日本のデータセンター建設（年間）", [3222, 5000, 10000],
     ["’23", "’24", "’28"], "1兆円超", "約3倍"),
]
ty, rh = 1.78, 1.3
for i, (what, note, vals, yrs, vlab, mult) in enumerate(rows):
    y = ty + i * rh
    if i > 0:
        hrule(s, LM, y-0.03, CW)
    text(s, LM, y+0.28, 3.1, 0.85,
         [[R(what, 17, INK, True)], [R(note, 10.5, MUTE, True, FONT_M)]],
         space_after=4, line_spacing=1.16)
    vbars(s, 4.15, y+rh-0.42, 3.9, 0.66, vals, yrs, vlab)
    text(s, 8.65, y+0.2, 1.75, 0.85, [[R(mult, 38, RED, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.45, y+0.2, 2.0, 0.85, [[R("に伸びた", 13, MUTE, True)]], anchor=MSO_ANCHOR.MIDDLE)

takeaway(s, 5.95,
         [[R("投資のピークは2027〜2028年。この波は、まだ“", 19, INK, True),
           R("序盤", 19, RED, True), R("”だ。", 19, INK, True)]], h=0.62)
text(s, LM, 6.82, 11.6, 0.28,
     [[R("出典：NVIDIA IR／Big Tech各社IR・報道／IDC Japan（国内DC建設投資）。", 9, FAINT, False, FONT_M)]])

# =========================================================
# SLIDE 4 — バリューチェーン概観
# =========================================================
s = slide()
header(s, "03", "AI需要 → 連鎖して伸びる、日本のバリューチェーン",
       "AIが伸びれば、この“川”がまるごと潤う。")

text(s, LM, 1.82, 8.5, 0.34, [[R("AI・DC需要の拡大", 15, RED, True),
     R("　この一手が、川上から川下までまるごと波及する", 11, MUTE, True)]])
# 川上 → 川下 の軸ラベル
text(s, LM, 2.3, 4.0, 0.22, [[R("川上（源流）", 9.5, FAINT, True, FONT_M, 1)]])
text(s, RM-4.0, 2.3, 4.0, 0.22, [[R("川下（据付・保守の現場）", 9.5, FAINT, True, FONT_M, 1)]],
     align=PP_ALIGN.RIGHT)
# “川”の流れ＝赤の全幅バー＋下流への矢尻
box(s, LM, 2.56, CW-0.34, 0.05, RED)
text(s, RM-0.36, 2.45, 0.4, 0.26, [[R("▶", 14, RED, True)]])

stages = [
    ("01", "発電・電源", "電気をつくる", ["三菱重工", "川崎重工", "IHI", "デンヨー"]),
    ("02", "送変電・電線", "送る・変える", ["日立", "三菱電機", "ダイヘン", "フジクラ"]),
    ("03", "DC建設・設備", "箱を建てる", ["鹿島建設", "大林組", "きんでん", "高砂熱学"]),
    ("04", "冷却・電源保護", "冷やす・止めない", ["ダイキン", "GSユアサ", "荏原製作所"]),
    ("05", "半導体", "計算する頭脳", ["東京エレクトロン", "ディスコ", "キオクシア", "信越化学"]),
]
spitch = CW / 5
stop = 2.92
for i, (no, stg, fn, cos) in enumerate(stages):
    x = LM + i * spitch
    w = spitch - 0.28
    if i > 0:
        vrule(s, x-0.12, stop, 2.6, wt=1.0)
    box(s, x, stop, w, 0.06, RED)
    text(s, x, stop+0.16, w, 0.22, [[R(no, 10, FAINT, True, FONT_M, 1)]])
    text(s, x, stop+0.42, w, 0.4, [[R(stg, 14, INK, True)]], line_spacing=1.05)
    text(s, x, stop+0.86, w, 0.24, [[R(fn, 10.5, MUTE, True)]])
    for j, co in enumerate(cos):
        text(s, x, stop+1.26 + j*0.40, w, 0.34,
             [[R("・", 11, RED, True), R(co, 12.5, SUB, True)]])

takeaway(s, 5.74,
         [[R("AIの“源流”が、川下の日本メーカーまで、", 19, INK, True),
           R("まるごと潤す。", 19, RED, True)]], h=0.56)
text(s, LM+0.28, 6.4, 11.3, 0.34,
     [[R("発電から半導体まで全段が同時に増収増益——狙える現場は、この一社一社にある。", 14, SUB, True)]])
text(s, LM, 6.94, 11.6, 0.28,
     [[R("※ 各段階の代表企業を抜粋（社名表記＝ロゴのイメージ、実ロゴへ差し替え可）。数値・詳細は次頁以降。", 9, FAINT, False, FONT_M)]])

# =========================================================
# SLIDE 5 — 建設費の中身：ANDPADの市場はどこか（赤で強弱）
# =========================================================
s = slide()
header(s, "04", "建設費の中身：どこにANDPADの市場があるか",
       "お金の3/4は設備。効くのは機器代ではなく“工”＝据付・試運転・保守。")

# Bar A：建設費 100%  ── 設備(75%)=赤、建設(25%)=グレー
bx, by, bw, bh = LM, 2.2, CW, 0.98
wg = bw*0.25; we = bw*0.50; wm = bw*0.25
box(s, bx, by, wg, bh, G_DK)
box(s, bx+wg, by, we, bh, RED)
box(s, bx+wg+we, by, wm, bh, RED2)
box(s, bx+wg-0.012, by, 0.024, bh, WHITE)
box(s, bx+wg+we-0.012, by, 0.024, bh, WHITE)
text(s, bx, by, wg, bh, [[R("建築（躯体）", 11, WHITE, True)], [R("約25%", 19, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=1, line_spacing=1.0)
text(s, bx+wg, by, we, bh, [[R("電気設備", 12, WHITE, True)], [R("約50%", 22, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=1, line_spacing=1.0)
text(s, bx+wg+we, by, wm, bh, [[R("空調・機械", 11, WHITE, True)], [R("約25%", 19, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=1, line_spacing=1.0)
text(s, bx, by-0.32, wg, 0.26, [[R("建設（ゼネコン）", 10.5, MUTE, True, FONT_M)]])
text(s, bx+wg, by-0.32, we+wm, 0.26, [[R("設備＝製造業（電気・機械）  約75%", 11, RED, True, FONT_M)]])

# Bar B：材 vs 工
text(s, LM, 3.45, CW, 0.28,
     [[R("その「設備 約75%」を  ", 12.5, SUB, True),
       R("材（機器・材料）", 12.5, MUTE, True), R("  と  ", 12.5, SUB),
       R("工（据付・施工＝人が動く）", 12.5, RED, True), R("  に分けると", 12.5, SUB, True)]])
b2y, b2h = 3.82, 0.88
wmat = CW*0.60; wwork = CW*0.40
box(s, LM, b2y, wmat, b2h, G1)
box(s, LM+wmat, b2y, wwork, b2h, RED)
box(s, LM+wmat-0.012, b2y, 0.024, b2h, WHITE)
text(s, LM, b2y, wmat, b2h, [[R("材：機器・材料（変圧器・冷凍機・UPS 等）  約6割", 12.5, INK, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, LM+wmat, b2y, wwork, b2h, [[R("工：据付・施工  約4割", 12.5, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, LM+wmat, b2y+b2h+0.06, wwork, 0.24,
     [[R("↑ ここがANDPADの市場（工）", 10, RED, True, FONT_M)]], align=PP_ALIGN.CENTER)

# 工の3工程
wy = 5.28
work = [
    ("据付・施工", "建設費の約3割", "国内の設備工事そのもの＝中核市場"),
    ("試運転", "建設費の1〜3%", "世界の試運転市場 $2.1B → $4.3B"),
    ("保守・運用", "毎年・運用費の約4割", "世界のDC運用保守 $15.8B → $45.6B"),
]
wpitch = CW / 3
for i, (nm, pct, mk) in enumerate(work):
    x = LM + i*wpitch
    if i > 0:
        vrule(s, x-0.2, wy, 0.7, wt=1.2)
    text(s, x, wy, wpitch-0.4, 0.3, [[R(nm, 13.5, INK, True), R("  "+pct, 12, RED, True)]])
    text(s, x, wy+0.36, wpitch-0.4, 0.32, [[R(mk, 10.5, MUTE, True)]], line_spacing=1.14)

takeaway(s, 6.32,
         [[R("建設業だけを見てきた。その周辺の“工”（据付・試運転・保守）に、", 16, INK, True),
           R("広大な市場", 16, RED, True), R("がある。", 16, INK, True)]], h=0.56)

# =========================================================
# SLIDE 6 — 主要企業：国内のどこに効いているか（裏付け）
# =========================================================
s = slide()
header(s, "05", "バリューチェーン別・国内主要プレーヤーと伸び",
       "国内のAI投資は、製造業のこの現場に着地している。")

flow = [
    ("01", "発電・電源", "電気をつくる", "受注残 5兆円超", "三菱重工"),
    ("02", "送変電・電線", "電気を送る・変える", "変圧器 生産能力2倍", "日立・ダイヘン・フジクラ"),
    ("03", "DC建設・設備工事", "箱を建てる", "空調工事 受注 +25.5%", "ダイダン・高砂熱学"),
    ("04", "冷却・電源保護", "冷やす・止めない", "北米DC冷却 約13倍", "ダイキン・GSユアサ"),
    ("05", "半導体", "計算する頭脳", "設備投資 +66%", "キオクシア・東京ｴﾚｸﾄﾛﾝ"),
]
fpitch = CW / 5
ftop = 2.0
for i, (no, stg, sb, stat, cos) in enumerate(flow):
    x = LM + i * fpitch
    if i > 0:
        vrule(s, x-0.12, ftop, 3.0, wt=1.2)
    text(s, x, ftop, fpitch-0.28, 0.22, [[R(no, 10, FAINT, True, FONT_M, 1)]])
    text(s, x, ftop+0.28, fpitch-0.24, 0.5, [[R(stg, 14, INK, True)]], line_spacing=1.05)
    text(s, x, ftop+0.82, fpitch-0.24, 0.26, [[R(sb, 10.5, MUTE, True)]])
    text(s, x, ftop+1.28, fpitch-0.3, 0.78, [[R(stat, 16, RED, True)]], line_spacing=1.1)
    text(s, x, ftop+2.35, fpitch-0.26, 0.6, [[R(cos, 10.5, SUB, True)]], line_spacing=1.18)

hrule(s, LM, 5.2, CW)
text(s, LM, 5.42, 11.6, 0.5,
     [[R("我々の入り方：", 12, MUTE, True, FONT_M),
       R("  A 発注者＝“建てる側”で管理", 13, SUB, True),
       R("   /   ", 12, FAINT),
       R("B 請負＝据付・試運転・保守を“納める側”で管理（本命）", 13, RED, True)]])
takeaway(s, 6.15,
         [[R("川上から川下まで、全段が同時に増収増益。", 17, INK, True),
           R("狙える現場が、日本中に生まれている。", 17, RED, True)]], h=0.6)

# =========================================================
# SLIDE 7 — クロージング
# =========================================================
s = slide()
box(s, LM, 1.2, 0.15, 0.15, RED)
text(s, 1.16, 1.12, 10.5, 0.3, [[R("SO, LET'S GO", 13, RED, True, FONT_M, 2)]])
text(s, LM-0.04, 1.78, 11.8, 1.2,
     [[R("次の主戦場は、", 44, INK, True), R("製造業", 44, RED, True), R("だ。", 44, INK, True)]])
box(s, LM, 3.35, 3.2, 0.05, RED)

pts = [
    ("市場は建設の3倍広い", "DC建設費の約3/4は電気・機械設備＝製造業。据付・試運転・保守まで現場が続く。"),
    ("追い風は本物", "AIチップ需要は3年で約13倍、投資ピークは2027-28。受注残に裏打ちされた構造需要。"),
    ("武器は完成済み", "建設で磨いた現場管理は、製造業の据付・保守フィールドに“そのまま効く”。"),
]
py = 3.85
for i, (h, b) in enumerate(pts):
    y = py + i * 0.92
    if i > 0:
        hrule(s, LM, y-0.08, CW)
    text(s, LM, y, 0.6, 0.65, [[R(f"0{i+1}", 17, RED, True, FONT_M)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, LM+0.75, y+0.06, 4.5, 0.6, [[R(h, 17, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, LM+5.4, y+0.06, 6.2, 0.62, [[R(b, 11.5, SUB)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.16)

box(s, LM, 6.72, CW, 0.05, RED)
text(s, LM-0.02, 6.9, 11.7, 0.5,
     [[R("建設の“周辺”に広がる巨大市場を、100人でつかみにいこう。", 21, INK, True)]])

# =========================================================
# APPENDIX — バリューチェーン別・キー企業一覧（全社網羅）
# =========================================================
DATA = Path(__file__).resolve().parent.parent / "data"

with open(DATA / "segments.csv", encoding="utf-8") as f:
    SEGMENTS = sorted(csv.DictReader(f), key=lambda r: int(r["value_chain_order"]))
with open(DATA / "companies.csv", encoding="utf-8") as f:
    _COMPS = list(csv.DictReader(f))
BY_SEG = {}
for _c in _COMPS:
    BY_SEG.setdefault(_c["segment_id"], []).append(_c)
TOTAL = sum(len(v) for v in BY_SEG.values())

# 表示用の簡潔なセグメント名（本編バリューチェーンと対応）
SEG_TITLE = {
    "SEG001": "特別高圧・大型変圧器",       "SEG002": "発電機・タービン・非常用電源",
    "SEG003": "送配電（開閉装置・遮断器）",  "SEG004": "DC建設ゼネコン",
    "SEG005": "DC電気・空調サブコン",       "SEG006": "産業冷却・空調（液冷・チラー）",
    "SEG007": "電線・ケーブル・光ファイバー", "SEG008": "半導体前工程（ファウンドリ/メモリ）",
    "SEG009": "半導体後工程（パッケージング）", "SEG010": "半導体製造装置・部材",
    "SEG011": "電力インフラ（発送電・系統増強）", "SEG012": "非常用発電機・UPS",
    "SEG013": "防音・免震・セキュリティ",     "SEG014": "建設資材・特殊建材",
    "SEG015": "半導体装置 部品・サブシステム", "SEG016": "電子部品・基板・産業ガス",
    "SEG017": "DC構築SI・DC運営",
}
CONF_COLOR = {"高": RED, "中": G_DK, "低": FAINT}
CC_SHORT = {"米国": "米", "アメリカ": "米", "台湾": "台", "韓国": "韓", "中国": "中",
            "オランダ": "蘭", "ドイツ": "独", "スイス": "瑞", "フランス": "仏",
            "英国": "英", "イギリス": "英", "アイルランド": "愛", "フィンランド": "芬",
            "スウェーデン": "典", "デンマーク": "丁", "ノルウェー": "諾", "カナダ": "加",
            "イタリア": "伊"}


def conf_c(v):
    return CONF_COLOR.get(v, FAINT)


def short_name(nm):
    for br in ("（", "("):
        i = nm.find(br)
        if i > 4:
            return nm[:i].strip()
    return nm


def norm_country(ct):
    ct = (ct or "").strip()
    for sep in "／/（(、 ":            # "ドイツ／英国（…）" → "ドイツ"
        p = ct.find(sep)
        if p > 0:
            ct = ct[:p]
    return ct.strip()


def ticker_label(c):
    tk = (c.get("ticker") or "").strip()
    if not tk:
        return "非上場"
    ct = norm_country(c.get("country"))
    if ct and ct != "日本":
        return f"{CC_SHORT.get(ct, ct[:1])} {tk}"
    return tk


# ---- 付録・扉 ----
s = slide()
box(s, LM, 1.35, 0.15, 0.15, RED)
text(s, 1.16, 1.27, 10, 0.3, [[R("APPENDIX", 13, RED, True, FONT_M, 2)]])
text(s, LM-0.04, 1.95, 11.8, 1.0,
     [[R("バリューチェーン別 ・ ", 34, INK, True), R("キー企業一覧", 34, RED, True)]])
box(s, LM, 3.15, 3.2, 0.05, RED)
text(s, LM-0.02, 3.45, 11.7, 0.5,
     [[R(f"AI需要の拡大で増収増益が見込める国内外 {TOTAL} 社を、17セグメント・バリューチェーン順に整理。",
         15, SUB)]])
ly = 4.35
text(s, LM, ly, 5, 0.28, [[R("確度（一次情報の裏付け）", 10.5, MUTE, True, FONT_M, 1)]])
for i, (lab, desc) in enumerate([("高", "受注・投資計画など一次情報が明確"),
                                 ("中", "一次情報＋報道等の二次情報"),
                                 ("低", "二次情報が中心")]):
    yy = ly + 0.42 + i * 0.42
    box(s, LM, yy+0.03, 0.15, 0.15, conf_c(lab))
    text(s, LM+0.3, yy, 6.5, 0.3, [[R(lab, 11, INK, True), R("　"+desc, 11, SUB)]])
text(s, 7.1, ly, 5, 0.28, [[R("社名右の表記", 10.5, MUTE, True, FONT_M, 1)]])
for i, (lab, desc) in enumerate([("6501 等", "証券コード（上場）"),
                                 ("非上場", "未上場・子会社等"),
                                 ("米 GEV 等", "海外は国名を併記")]):
    yy = ly + 0.42 + i * 0.42
    text(s, 7.1, yy, 1.5, 0.3, [[R(lab, 10.5, FAINT, True, FONT_M)]])
    text(s, 8.7, yy, 3.7, 0.3, [[R(desc, 11, SUB)]])
box(s, LM, 6.7, CW, 0.03, HAIR)
text(s, LM, 6.84, 11.7, 0.3,
     [[R("出典：各社決算・IR等の一次情報＋公開データ。詳細は企業マスタ valuechain_master.xlsx ／ research/ を参照。",
         9.5, FAINT, False, FONT_M)]])


def appendix_header(s, page):
    box(s, LM, 0.6, 0.14, 0.14, RED)
    text(s, 1.14, 0.53, 9.5, 0.28,
         [[R("APPENDIX ・ バリューチェーン別 キー企業一覧", 11.5, RED, True, FONT_M, 1.5)]])
    text(s, RM-1.6, 0.53, 1.6, 0.28, [[R(page, 11.5, FAINT, True, FONT_M)]], align=PP_ALIGN.RIGHT)
    hrule(s, LM, 0.96, CW, wt=1.2)


# ---- 2カラム・フロー配置で全社を流し込む ----
COL_GAP = 0.6
COL_W = (CW - COL_GAP) / 2
COL_X = [LM, LM + COL_W + COL_GAP]
A_TOP, A_BOT, ROW_H, HDR_H = 1.18, 7.15, 0.25, 0.46

stream = []
for seg in SEGMENTS:
    cos = BY_SEG.get(seg["segment_id"], [])
    if cos:
        stream.append(("seg", seg))
        stream.extend(("co", c) for c in cos)

pages, cur, col, y = [], [], 0, A_TOP
for kind, data in stream:
    h = HDR_H if kind == "seg" else ROW_H
    need = h + (2 * ROW_H if kind == "seg" else 0)   # ヘッダーの孤立を防ぐ
    if y + need > A_BOT:
        if col == 0:
            col, y = 1, A_TOP
        else:
            pages.append(cur); cur, col, y = [], 0, A_TOP
    cur.append((col, y, kind, data))
    y += h
if cur:
    pages.append(cur)

for pi, cmds in enumerate(pages):
    s = slide()
    appendix_header(s, f"付録 {pi+1} / {len(pages)}")
    for col, y, kind, data in cmds:
        x = COL_X[col]
        if kind == "seg":
            box(s, x, y+0.07, 0.12, 0.25, RED)
            title = SEG_TITLE.get(data["segment_id"], data["segment_name"])
            n = len(BY_SEG.get(data["segment_id"], []))
            text(s, x+0.24, y, COL_W-0.24, 0.34,
                 [[R(f"{int(data['value_chain_order']):02d}  ", 10, RED, True, FONT_M),
                   R(title, 11.5, INK, True),
                   R(f"　{n}社", 9, FAINT, True, FONT_M)]], line_spacing=1.0)
            hrule(s, x, y+0.4, COL_W, color=HAIR, wt=0.8)
        else:
            box(s, x, y+0.06, 0.1, 0.1, conf_c(data.get("confidence", "")))
            text(s, x+0.22, y-0.02, COL_W-1.62, 0.26,
                 [[R(short_name(data["company_name_ja"]), 10.5, SUB, True)]])
            text(s, x+COL_W-1.35, y-0.02, 1.35, 0.26,
                 [[R(ticker_label(data), 9, FAINT, True, FONT_M)]], align=PP_ALIGN.RIGHT)

out = Path(__file__).resolve().parent.parent / "output" / "andpad_soukai_ai_deck.pptx"
prs.save(str(out))
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst),
      "/ appendix pages:", len(pages), "/ companies:", TOTAL)
